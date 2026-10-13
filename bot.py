import os, re, logging, shutil, zipfile
from pathlib import Path
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import Application, CommandHandler, MessageHandler, CallbackQueryHandler, ContextTypes, filters
from ebooklib import epub, ITEM_DOCUMENT
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader
from deep_translator import GoogleTranslator

logging.basicConfig(level=logging.INFO)
TOKEN = os.getenv("BOT_TOKEN", "").strip()
WORK = Path(os.getenv("WORKDIR", "work"))
WORK.mkdir(parents=True, exist_ok=True)
MAX_DOWNLOAD = int(os.getenv("MAX_DOWNLOAD_MB", "20")) * 1024 * 1024
JOBS = {}
SUPPORTED = {".txt", ".epub", ".docx", ".pdf", ".html", ".htm"}

def safe_name(name):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", name or "file")[:100] or "file"

def read_file(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".txt":
        return path.read_text(encoding="utf-8-sig", errors="replace")
    if ext == ".docx":
        doc = Document(str(path))
        blocks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                blocks.append(" | ".join(c.text for c in row.cells))
        return "\n\n".join(blocks)
    if ext == ".epub":
        book = epub.read_epub(str(path))
        chunks = []
        for item in book.get_items():
            if item.get_type() == ITEM_DOCUMENT:
                soup = BeautifulSoup(item.get_content(), "html.parser")
                for tag in soup(["script", "style"]): tag.decompose()
                txt = soup.get_text("\n", strip=True)
                if txt: chunks.append(txt)
        return "\n\n".join(chunks)
    if ext == ".pdf":
        reader = PdfReader(str(path))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if ext in (".html", ".htm"):
        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
        for tag in soup(["script", "style"]): tag.decompose()
        return soup.get_text("\n", strip=True)
    raise ValueError("Unsupported input format.")

def write_txt(text, path):
    path.write_text(text, encoding="utf-8")

def write_docx(text, path):
    doc = Document()
    for block in re.split(r"\n\s*\n", text):
        if block.strip():
            doc.add_paragraph(block.strip())
    doc.save(str(path))

def write_epub(text, path, title="Converted Novel"):
    book = epub.EpubBook()
    book.set_identifier("converted-" + str(abs(hash(text[:2000]))))
    book.set_title(title)
    book.set_language("en")
    chapter = epub.EpubHtml(title="Novel", file_name="chapter_001.xhtml", lang="en")
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    html = "".join("<p>" + p.replace("&","&amp;").replace("<","&lt;").replace(">","&gt;").replace("\n","<br/>") + "</p>" for p in paras)
    chapter.content = ("<html><head><title>" + title + "</title></head><body>" + html + "</body></html>").encode("utf-8")
    book.add_item(chapter)
    book.toc = (epub.Link("chapter_001.xhtml", "Novel", "novel"),)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", chapter]
    epub.write_epub(str(path), book)

def write_html(text, path, title="Converted file"):
    import html
    paras = "".join("<p>" + html.escape(p).replace("\n", "<br>") + "</p>" for p in re.split(r"\n\s*\n", text) if p.strip())
    path.write_text(f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title></head><body>{paras}</body></html>", encoding="utf-8")

def split_text(text, mode, value):
    # Split by chapter headings if possible, otherwise by requested character chunks.
    if mode == "chapter":
        pattern = re.compile(r"(?im)^(?=(?:chapter|chap\.?|第.{1,12}章|卷[一二三四五六七八九十百0-9]+)[^\n]{0,100}$)")
        parts = [p.strip() for p in pattern.split(text) if p.strip()]
        if len(parts) > 1:
            return parts
        mode = "chars"
        value = max(10000, int(value) * 5000)
    size = max(1000, int(value))
    return [text[i:i+size].strip() for i in range(0, len(text), size) if text[i:i+size].strip()]

def translate_text(text, target):
    # Translation service has per-request limits; translate paragraph/sentence chunks.
    translated = []
    translator = GoogleTranslator(source="auto", target=target)
    chunks = re.split(r"(\n\s*\n)", text)
    for chunk in chunks:
        if not chunk.strip():
            translated.append(chunk)
            continue
        # Keep requests comfortably below common provider limits.
        for i in range(0, len(chunk), 3500):
            translated.append(translator.translate(chunk[i:i+3500]) or chunk[i:i+3500])
    return "".join(translated)

def output_file(src: Path, fmt: str, text: str, destdir: Path):
    dest = destdir / (src.stem + "." + fmt)
    if fmt == "txt": write_txt(text, dest)
    elif fmt == "docx": write_docx(text, dest)
    elif fmt == "epub": write_epub(text, dest, src.stem)
    elif fmt == "html": write_html(text, dest, src.stem)
    else: raise ValueError("That output format is not supported.")
    return dest

def keyboard():
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("Convert", callback_data="convert"), InlineKeyboardButton("Split", callback_data="split")],
        [InlineKeyboardButton("Chinese → English", callback_data="tr_en"), InlineKeyboardButton("Chinese → Hindi", callback_data="tr_hi")],
        [InlineKeyboardButton("English → Hindi", callback_data="en_hi"), InlineKeyboardButton("Hindi → English", callback_data="hi_en")],
        [InlineKeyboardButton("Help", callback_data="help")]
    ])

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "📚 Novel Convert + Split + Translate Bot\n\n"
        "Send a TXT, EPUB, DOCX, PDF, or HTML file, then choose an action.\n\n"
        "Priority: EPUB → DOCX. Also supports TXT ↔ DOCX/EPUB/HTML where offered.\n"
        "Translation uses an online service and needs internet; it is not offline.\n"
        "Use /help for commands.", reply_markup=keyboard())

async def help_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "How to use:\n1) Send a supported file.\n2) Tap Convert, Split, or a translation button.\n\n"
        "Commands:\n/convert docx — convert latest file to DOCX\n/convert txt — convert latest file to TXT\n/convert epub — convert latest file to EPUB\n"
        "/split 20000 — split into ~20,000-character parts\n/split chapter — split by detected chapter headings\n"
        "/translate en — translate latest file to English\n/translate hi — translate latest file to Hindi\n\n"
        "EPUB → DOCX extracts readable chapter text and does not preserve original images or exact formatting. Scanned PDFs may have no extractable text.")

async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    msg = update.message
    doc = msg.document
    ext = Path(doc.file_name or "").suffix.lower()
    if ext not in SUPPORTED:
        await msg.reply_text("Unsupported format. Send TXT, EPUB, DOCX, PDF, HTML, or HTM.")
        return
    if doc.file_size and doc.file_size > MAX_DOWNLOAD:
        await msg.reply_text(f"File too large for this bot setting. Limit: {MAX_DOWNLOAD // (1024*1024)} MB.")
        return
    uid = msg.from_user.id
    jobdir = WORK / str(uid)
    shutil.rmtree(jobdir, ignore_errors=True)
    jobdir.mkdir(parents=True, exist_ok=True)
    path = jobdir / safe_name(doc.file_name)
    try:
        tgfile = await context.bot.get_file(doc.file_id)
        await tgfile.download_to_drive(custom_path=str(path))
        text = read_file(path)
        if not text.strip():
            await msg.reply_text("I couldn't extract text from this file. Scanned PDFs need OCR, which isn't included.")
            return
        JOBS[uid] = {"path": path, "text": text, "jobdir": jobdir}
        await msg.reply_text(f"✅ Received: {doc.file_name}\nSize: {len(text):,} characters\nChoose an action:", reply_markup=keyboard())
    except Exception as e:
        logging.exception("file handling failed")
        await msg.reply_text(f"❌ Could not read this file: {str(e)[:250]}")

async def send_path(update, path: Path, caption="Done"):
    with path.open("rb") as f:
        await update.effective_message.reply_document(document=f, filename=path.name, caption=caption[:900])

async def do_convert(update, context, fmt):
    uid = update.effective_user.id
    job = JOBS.get(uid)
    if not job: return await update.effective_message.reply_text("Send a file first.")
    try:
        dest = output_file(job["path"], fmt, job["text"], job["jobdir"])
        await send_path(update, dest, f"✅ Converted to {fmt.upper()}")
    except Exception as e:
        logging.exception("convert failed")
        await update.effective_message.reply_text(f"❌ Conversion failed: {str(e)[:250]}")

async def do_split(update, context, mode="chars", value="20000"):
    uid = update.effective_user.id
    job = JOBS.get(uid)
    if not job: return await update.effective_message.reply_text("Send a file first.")
    try:
        parts = split_text(job["text"], mode, value)
        outdir = job["jobdir"] / "split_parts"
        shutil.rmtree(outdir, ignore_errors=True); outdir.mkdir()
        for i, part in enumerate(parts, 1):
            write_txt(part, outdir / f"part_{i:04}.txt")
        archive = job["jobdir"] / f"{job['path'].stem}_split.zip"
        with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as z:
            for p in outdir.iterdir(): z.write(p, p.name)
        await send_path(update, archive, f"✂️ Split into {len(parts)} TXT parts.")
    except Exception as e:
        logging.exception("split failed")
        await update.effective_message.reply_text(f"❌ Split failed: {str(e)[:250]}")

async def do_translate(update, context, target):
    uid = update.effective_user.id
    job = JOBS.get(uid)
    if not job: return await update.effective_message.reply_text("Send a file first.")
    status = await update.effective_message.reply_text("🌐 Translating. Large novels may take a while…")
    try:
        translated = translate_text(job["text"], target)
        fmt = "docx" if job["path"].suffix.lower() in (".epub", ".docx", ".pdf") else "txt"
        dest = output_file(job["path"], fmt, translated, job["jobdir"])
        dest = dest.with_name(job["path"].stem + f"_{target}." + fmt)
        # output_file already wrote; rename safely
        original = job["jobdir"] / (job["path"].stem + "." + fmt)
        if original.exists() and original != dest: original.replace(dest)
        await status.delete()
        await send_path(update, dest, f"🌐 Translation complete: {target.upper()} ({fmt.upper()})")
    except Exception as e:
        logging.exception("translation failed")
        await status.edit_text(f"❌ Translation failed: {str(e)[:250]}\nSome free translation services may rate-limit large jobs.")

async def button_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    uid = q.from_user.id
    job = JOBS.get(uid)
    if q.data == "help":
        return await q.message.reply_text("Send a supported file, then tap Convert / Split / translation. EPUB → DOCX is available via Convert.")
    if not job:
        return await q.message.reply_text("Send a file first.")
    if q.data == "convert":
        ext = job["path"].suffix.lower()
        options = ["docx", "txt", "epub", "html"]
        await q.message.reply_text("Choose output format:", reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton(x.upper(), callback_data="fmt_"+x) for x in options]]))
    elif q.data.startswith("fmt_"):
        await do_convert(q, context, q.data[4:])
    elif q.data == "split":
        await q.message.reply_text("Split mode:", reply_markup=InlineKeyboardMarkup([
            [InlineKeyboardButton("By chapter headings", callback_data="split_chapter")],
            [InlineKeyboardButton("By 20,000 characters", callback_data="split_chars_20000")],
            [InlineKeyboardButton("By 50,000 characters", callback_data="split_chars_50000")]
        ]))
    elif q.data == "split_chapter":
        await do_split(q, context, "chapter", "1")
    elif q.data.startswith("split_chars_"):
        await do_split(q, context, "chars", q.data.rsplit("_",1)[1])
    elif q.data == "tr_en": await do_translate(q, context, "english")
    elif q.data == "tr_hi": await do_translate(q, context, "hindi")
    elif q.data == "en_hi": await do_translate(q, context, "hindi")
    elif q.data == "hi_en": await do_translate(q, context, "english")

async def convert_cmd(update, context):
    if not context.args: return await update.message.reply_text("Usage: /convert docx")
    await do_convert(update, context, context.args[0].lower().lstrip("."))

async def split_cmd(update, context):
    mode = context.args[0].lower() if context.args else "20000"
    if mode == "chapter": await do_split(update, context, "chapter", "1")
    else:
        try: size = str(max(1000, min(100000, int(mode))))
        except ValueError: size = "20000"
        await do_split(update, context, "chars", size)

async def translate_cmd(update, context):
    if not context.args: return await update.message.reply_text("Usage: /translate en  OR  /translate hi")
    lang = context.args[0].lower()
    if lang in ("en", "english"): await do_translate(update, context, "english")
    elif lang in ("hi", "hindi"): await do_translate(update, context, "hindi")
    else: await update.message.reply_text("Supported targets: en, hi")

def main():
    if not TOKEN:
        raise SystemExit("BOT_TOKEN environment variable is missing.")
    app = Application.builder().token(TOKEN).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("help", help_cmd))
    app.add_handler(CommandHandler("convert", convert_cmd))
    app.add_handler(CommandHandler("split", split_cmd))
    app.add_handler(CommandHandler("translate", translate_cmd))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_document))
    app.add_handler(CallbackQueryHandler(button_handler))
    app.run_polling(drop_pending_updates=True)

if __name__ == "__main__":
    main()
